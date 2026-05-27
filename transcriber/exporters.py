from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING

from transcriber.utils import format_srt_timestamp, format_timestamp
if TYPE_CHECKING:
    from transcriber.whisper_engine import TranscriptSegment


def build_txt_content(segments: list["TranscriptSegment"]) -> str:
    return "\n".join(f"[{format_timestamp(seg.start)}] {seg.text}" for seg in segments).strip() + "\n"


def build_markdown_content(
    title: str,
    source_url: str,
    created_iso: str,
    model: str,
    language: str | None,
    segments: list["TranscriptSegment"],
) -> str:
    language_value = language or "auto"
    body = "\n".join(f"[{format_timestamp(seg.start)}] {seg.text}" for seg in segments)
    return (
        "---\n"
        f'title: "{title}"\n'
        f'source: "{source_url}"\n'
        f'created: "{created_iso}"\n'
        f'model: "{model}"\n'
        f'language: "{language_value}"\n'
        "---\n\n"
        f"# {title}\n\n"
        f"Source: {source_url}\n\n"
        "## Transcript\n\n"
        f"{body}\n"
    )


def build_srt_content(segments: list["TranscriptSegment"]) -> str:
    lines: list[str] = []
    for idx, seg in enumerate(segments, start=1):
        lines.extend(
            [
                str(idx),
                f"{format_srt_timestamp(seg.start)} --> {format_srt_timestamp(seg.end)}",
                seg.text,
                "",
            ]
        )
    return "\n".join(lines).rstrip() + "\n"


def export_txt(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8")


def export_markdown(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8")


def export_srt(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8")


def export_docx(
    path: Path,
    title: str,
    source_url: str,
    created_iso: str,
    model: str,
    language: str | None,
    segments: list["TranscriptSegment"],
) -> None:
    from docx import Document

    document = Document()
    document.add_heading(title, level=1)
    document.add_paragraph(f"Source URL: {source_url}")
    document.add_paragraph(f"Created: {created_iso}")
    document.add_paragraph(f"Model: {model}")
    document.add_paragraph(f"Language: {language or 'auto'}")
    document.add_heading("Transcript", level=2)
    for seg in segments:
        document.add_paragraph(f"[{format_timestamp(seg.start)}] {seg.text}")
    document.save(path)
